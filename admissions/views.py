from django.db import transaction
from django.contrib.auth import get_user_model
from django.shortcuts import render
from openpyxl import load_workbook
from .models import Student
from .forms import StudentImportForm
User = get_user_model()

from academics.models import (
    Programme,
    ProgrammeEntryType,
    AdmissionCategory,
)


def import_students_view(request):

    rows = []
    errors = []

    if request.method == "POST":

        form = StudentImportForm(request.POST, request.FILES)

        if form.is_valid():

            excel_file = request.FILES["excel_file"]

            workbook = load_workbook(excel_file)

            sheet = workbook.active

            seen_roll_numbers = set()
            seen_aadhaar_numbers = set()

            for row_number, row in enumerate(
                sheet.iter_rows(values_only=True),
                start=1,
            ):

                rows.append(row)

                # Skip header row
                if row_number == 1:
                    continue

                programme = row[10]
                entry_type = row[11]
                category = row[12]
                roll_number = row[0]
                if roll_number:

                    if roll_number in seen_roll_numbers:
                        errors.append(
                            f"Row {row_number}: Duplicate Roll Number '{roll_number}' found in Excel."
                        )
                    else:
                        seen_roll_numbers.add(roll_number)
                student_name = row[1]
                father_name = row[2]
                mother_name = row[3]
                dob = row[4]
                aadhar = str(row[5]).strip() if row[5] is not None else ""
                if aadhar:

                    if aadhar in seen_aadhaar_numbers:
                        errors.append(
                            f"Row {row_number}: Duplicate Aadhaar '{aadhar}' found in Excel."
                        )
                    else:
                        seen_aadhaar_numbers.add(aadhar)
                
                mobile = row[6]
                parent_mobile = row[7]
                gender = row[8]
                blood_group = row[9]
                admission_year = row[13]
                email = row[14]
                remarks = row[15]
                if not roll_number:
                    errors.append(f"Row {row_number}: Roll Number is required.")

                if not student_name:
                    errors.append(f"Row {row_number}: Student Name is required.")

                if not father_name:
                    errors.append(f"Row {row_number}: Father Name is required.")

                if not mother_name:
                    errors.append(f"Row {row_number}: Mother Name is required.")

                if not dob:
                    errors.append(f"Row {row_number}: Date of Birth is required.")

                if not mobile:
                    errors.append(f"Row {row_number}: Mobile Number is required.")

                if not parent_mobile:
                    errors.append(f"Row {row_number}: Parent Mobile Number is required.")

                if not blood_group:
                    errors.append(f"Row {row_number}: Blood Group is required.")

                if not admission_year:
                    errors.append(f"Row {row_number}: Admission Academic Year is required.")

                # Validate Gender
                if gender and gender not in ["MALE", "FEMALE", "OTHER"]:
                    errors.append(
                        f"Row {row_number}: Gender must be MALE, FEMALE or OTHER."
                    )

                if not programme:
                    errors.append(f"Row {row_number}: Programme is required.")

                if not entry_type:
                    errors.append(f"Row {row_number}: Programme Entry Type is required.")

                if not category:
                    errors.append(f"Row {row_number}: Admission Category is required.")

                # Duplicate Roll Number in Database
                if roll_number and Student.objects.filter(
                    roll_number=roll_number
                ).exists():
                    errors.append(
                        f"Row {row_number}: Roll Number '{roll_number}' already exists in the database."
                    )
                
                # Duplicate Aadhaar in Database
                                
                if Student.objects.filter(
                    aadhaar_number=aadhar
                ).exists():
                    errors.append(
                        f"Row {row_number}: Aadhaar '{aadhar}' already exists in the database."
                    )

                if not Programme.objects.filter(
                    programme_name=programme
                ).exists():
                    errors.append(
                        f"Row {row_number}: Programme '{programme}' does not exist."
                    )

                if not ProgrammeEntryType.objects.filter(
                    entry_type_name=entry_type
                ).exists():
                    errors.append(
                        f"Row {row_number}: Programme Entry Type '{entry_type}' does not exist."
                    )

                if not AdmissionCategory.objects.filter(
                    category_name=category
                ).exists():
                    errors.append(
                        f"Row {row_number}: Admission Category '{category}' does not exist."
                    )
            if not errors:

                with transaction.atomic():

                    for row in rows[1:]:

                        roll_number = row[0]
                        student_name = row[1]
                        father_name = row[2]
                        mother_name = row[3]
                        dob = row[4]
                        aadhaar = str(row[5]).strip() if row[5] else ""
                        mobile = str(row[6]).strip() if row[6] else ""
                        parent_mobile = str(row[7]).strip() if row[7] else ""
                        gender = row[8]
                        blood_group = row[9]
                        programme_name = row[10]
                        entry_type_name = row[11]
                        category_name = row[12]
                        admission_year = row[13]
                        email = row[14]
                        remarks = row[15]

                        programme = Programme.objects.get(
                            programme_name=programme_name
                        )

                        entry_type = ProgrammeEntryType.objects.get(
                            entry_type_name=entry_type_name
                        )

                        category = AdmissionCategory.objects.get(
                            category_name=category_name
                        )

                        user = User.objects.create_user(
                            username=roll_number,
                            password=dob.strftime("%d%m%Y"),
                            first_name=student_name,
                            email=email or "",
                        )

                        Student.objects.create(
                            user=user,
                            roll_number=roll_number,
                            student_name=student_name,
                            father_name=father_name,
                            mother_name=mother_name,
                            aadhaar_number=aadhaar,
                            mobile_number=mobile,
                            parent_mobile_number=parent_mobile,
                            email=email,
                            date_of_birth=dob,
                            gender=gender,
                            blood_group=blood_group,
                            programme=programme,
                            programme_entry_type=entry_type,
                            admission_category=category,
                            admission_academic_year=admission_year,
                            remarks=remarks or "",
                            documents_verified=True,
                            admission_status="PENDING",
                            is_active=True,
                        )
    else:

        form = StudentImportForm()

    return render(
        request,
        "admin/admissions/import_students.html",
        {
            "form": form,
            "rows": rows,
            "errors": errors,
            "title": "Import Students",
        },
    )
