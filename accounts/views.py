import random
from decimal import Decimal
from results.services import (
    calculate_sgpa,
    calculate_cgpa,
    get_backlog_subjects,
    get_live_attempt_number,
    calculate_backlog_result,
    get_effective_results,
)
from academics.models import Subject
from openpyxl import load_workbook
from django.http import HttpResponse
from results.models import (
    ResultEntry,
    StudentAcademicResult,
    BacklogResultEntry,
    SemesterResult,
    CombinedResult,
    ResultPublication,
)
from django.db.models import Count
from examinations.models import (
    Examination,
    ExaminationSemester,
)
from evaluation.models import FacultySubjectAssignment
from django.db.models import Count
from django.contrib import messages
from django.utils import timezone
from examinations.models import (
    Examination,
    AssessmentWindow,
)
from results.models import ResultEntry
from django.shortcuts import redirect
from registration.models import ExaminationRegistration
from django.contrib.auth.decorators import login_required
from django.contrib.auth.views import LoginView
from django.shortcuts import render
from admissions.models import Student
from registration.models import (
    SemesterRegistration,
    BacklogRegistrationSubject,
)
from evaluation.models import (
    Faculty,
    FacultySubjectAssignment,
    BacklogFacultyAssignment,
    QuestionPaper,
    BacklogQuestionPaper,
)
from academics.models import (
    AcademicYear,
    Curriculum,
    Subject,
)
from evaluation.services import (
    submit_question_papers,
    submit_backlog_question_papers,
    generate_question_paper,
)
from registration.models import SemesterRegistration
from django.contrib.auth import logout
from django.shortcuts import redirect
from registration.models import (
    BacklogRegistration,
    BacklogRegistrationSubject,
)
from django.shortcuts import (
    get_object_or_404,
)

class LoginPageView(LoginView):

    template_name = "accounts/login.html"

    def get_success_url(self):

        user = self.request.user

        if user.role:

            if user.role.role_code == "STUDENT":
                return "/student/"

            if user.role.role_code == "ADMIN":
                return "/admin/"

            if user.role.role_code == "DEPM":
                return "/depm/"

            if user.role.role_code == "FACULTY":
                return "/faculty/"
    
            if user.role.role_code == "COE":
                return "/coe/"
        return "/login/"

@login_required
def student_dashboard(request):

    student = Student.objects.get(
        user=request.user
    )

    latest_registration = (
        SemesterRegistration.objects.filter(
            student=student,
        )
        .order_by("-registration_date")
        .first()
    )

    return render(
        request,
        "accounts/student_dashboard.html",
        {
            "student": student,
            "latest_registration": latest_registration,
        },
    )

@login_required
def depm_dashboard(request):

    return render(
        request,
        "accounts/depm_dashboard.html",
    )

@login_required
def faculty_subject_assignment(request):
    message = ""
    if request.method == "POST":

        
        academic_year = AcademicYear.objects.get(
            id=request.POST.get("academic_year")
        )

        curriculum = Curriculum.objects.get(
            id=request.POST.get("curriculum")
        )

        subject = Subject.objects.get(
            id=request.POST.get("subject")
        )

        faculty = Faculty.objects.get(
            id=request.POST.get("faculty")
        )

        assignment, created = FacultySubjectAssignment.objects.get_or_create(
            academic_year=academic_year,
            subject=subject,
            defaults={
                "curriculum": curriculum,
                "faculty": faculty,
            },
        )

        if created:
            message = "Faculty assigned successfully."
            print("MESSAGE:", message)
        else:
            message = "This subject is already assigned."
            print("MESSAGE:", message)

    return render(
        request,
        "accounts/faculty_subject_assignment.html",
        {
            "academic_years": AcademicYear.objects.filter(
                is_active=True,
            ),
            "curriculums": Curriculum.objects.filter(
                is_active=True,
            ),
            "subjects": Subject.objects.filter(
                is_active=True,
            ),
            "faculties": Faculty.objects.filter(
                is_active=True,
            ),
            "message": message,
        },
    )

@login_required
def backlog_faculty_assignment(request):

    message = ""

    examinations = Examination.objects.filter(
        examination_type__code="BLG",
        is_active=True,
    ).order_by("-id")

    selected_examination = None

    subjects = []

    subject_counts = {}

    faculties = Faculty.objects.filter(
        is_active=True,
    )

    existing_assignments = {}

    if request.method == "POST":

        examination_id = request.POST.get("examination")

        if examination_id:

            selected_examination = get_object_or_404(
                Examination,
                id=examination_id,
                examination_type__code="BLG",
                is_active=True,
            )

            action = request.POST.get("action")

            approved_subject_records = (
                BacklogRegistrationSubject.objects.filter(
                    backlog_registration__examination=selected_examination,
                    backlog_registration__registration_status="APPROVED",
                )
                .select_related(
                    "subject",
                    "backlog_registration",
                )
            )

            subject_map = {}
            subject_counts = {}

            for record in approved_subject_records:

                subject_map[record.subject.id] = record.subject

                subject_counts[record.subject.id] = (
                    subject_counts.get(record.subject.id, 0) + 1
                )

            subjects = sorted(
                subject_map.values(),
                key=lambda subject: (
                    subject.semester,
                    subject.display_order,
                ),
            )

            existing_assignments = {
                assignment.subject_id: assignment.faculty_id
                for assignment in BacklogFacultyAssignment.objects.filter(
                    examination=selected_examination,
                    is_active=True,
                )
            }

            for subject in subjects:

                subject.backlog_student_count = subject_counts.get(
                    subject.id,
                    0,
                )

                subject.assigned_faculty_id = existing_assignments.get(
                    subject.id
                )

            if action == "save":

                for subject in subjects:

                    faculty_id = request.POST.get(
                        f"faculty_{subject.id}"
                    )

                    if not faculty_id:
                        continue

                    faculty = get_object_or_404(
                        Faculty,
                        id=faculty_id,
                        is_active=True,
                    )

                    BacklogFacultyAssignment.objects.update_or_create(
                        examination=selected_examination,
                        subject=subject,
                        defaults={
                            "curriculum": subject.curriculum,
                            "faculty": faculty,
                            "is_active": True,
                        },
                    )

                message = (
                    "Backlog faculty assignments saved successfully."
                )

    return render(
        request,
        "accounts/backlog_faculty_assignment.html",
        {
            "examinations": examinations,
            "selected_examination": selected_examination,
            "subjects": subjects,
            "faculties": faculties,
            "message": message,
            "subject_counts": subject_counts,
            "existing_assignments": existing_assignments,
        },
    )

@login_required
def faculty_dashboard(request):

    faculty = Faculty.objects.get(
        user=request.user
    )

    assignments = FacultySubjectAssignment.objects.filter(
        faculty=faculty,
    )

    backlog_assignments = BacklogFacultyAssignment.objects.filter(
        faculty=faculty,
        is_active=True,
    ).select_related(
        "examination",
        "subject",
    )

    return render(
        request,
        "accounts/faculty_dashboard.html",
        {
            "faculty": faculty,
            "assignments": assignments,
            "backlog_assignments": backlog_assignments,
        },
    )

@login_required
def coe_dashboard(request):

    return render(
        request,
        "accounts/coe_dashboard.html",
    )

@login_required
def question_paper_selection(request):

    assignments = FacultySubjectAssignment.objects.filter(
        submitted_on__isnull=False,
    )

    question_papers = QuestionPaper.objects.all()

    return render(
        request,
        "accounts/question_paper_selection.html",
        {
            "assignments": assignments,
            "question_papers": question_papers,
        },
    )

@login_required
def backlog_question_paper_selection(request):

    assignments = BacklogFacultyAssignment.objects.filter(
        submitted_on__isnull=False,
        is_active=True,
    ).select_related(
        "examination",
        "subject",
        "faculty",
    )

    question_papers = BacklogQuestionPaper.objects.all()

    return render(
        request,
        "accounts/backlog_question_paper_selection.html",
        {
            "assignments": assignments,
            "question_papers": question_papers,
        },
    )

@login_required
def generate_backlog_question_paper_view(
    request,
    assignment_id,
):

    assignment = BacklogFacultyAssignment.objects.get(
        id=assignment_id,
        is_active=True,
    )

    if not assignment.selected_question_paper_set:

        assignment.selected_question_paper_set = random.choice(
            [1, 2, 3]
        )

        assignment.save(
            update_fields=[
                "selected_question_paper_set",
            ]
        )

    return redirect(
        "backlog_question_paper_selection",
    )

@login_required
def upload_question_papers(request, assignment_id):

    assignment = FacultySubjectAssignment.objects.get(
        id=assignment_id,
        faculty__user=request.user,
    )

    message = ""

    if request.method == "POST":

        try:

            submit_question_papers(

                assignment=assignment,

                set1=request.FILES.get("set1"),

                set2=request.FILES.get("set2"),

                set3=request.FILES.get("set3"),

                minutes=request.FILES.get("minutes"),

                declaration=(
                    request.POST.get("declaration")
                    == "on"
                ),
            )

            message = (
                "Confidential Question Papers "
                "submitted successfully."
            )

        except ValueError as e:

            message = str(e)

    return render(
        request,
        "accounts/upload_question_papers.html",
        {
            "assignment": assignment,
            "message": message,
        },
    )

@login_required
def upload_backlog_question_papers(request, assignment_id):

    assignment = get_object_or_404(
        BacklogFacultyAssignment,
        id=assignment_id,
        faculty__user=request.user,
        is_active=True,
    )

    message = ""

    if request.method == "POST":

        try:

            submit_backlog_question_papers(

                assignment=assignment,

                set1=request.FILES.get("set1"),

                set2=request.FILES.get("set2"),

                set3=request.FILES.get("set3"),

                minutes=request.FILES.get("minutes"),

                declaration=(
                    request.POST.get("declaration")
                    == "on"
                ),
            )

            message = (
                "Confidential Backlog Question Papers "
                "submitted successfully."
            )

        except ValueError as e:

            message = str(e)

    return render(
        request,
        "accounts/upload_backlog_question_papers.html",
        {
            "assignment": assignment,
            "message": message,
        },
    )

@login_required
def generate_question_paper_view(
    request,
    assignment_id,
):

    assignment = FacultySubjectAssignment.objects.get(
        id=assignment_id,
    )

    generate_question_paper(
        assignment,
    )

    return redirect(
        "question_paper_selection",
    )

@login_required
def enter_marks(request, assignment_id):

    assignment = FacultySubjectAssignment.objects.get(
        id=assignment_id,
        faculty__user=request.user,
    )

    examinations = Examination.objects.all().order_by(
        "-academic_year",
        "examination_name",
    )

    selected_examination_id = request.GET.get("examination")

    registrations = ExaminationRegistration.objects.none()

    selected_examination = None

    active_window = None

    if selected_examination_id:

        selected_examination = Examination.objects.get(
            id=selected_examination_id
        )

        active_window = (
            AssessmentWindow.objects.filter(
                examination=selected_examination,
                start_date__lte=timezone.localdate(),
                end_date__gte=timezone.localdate(),
            ).first()
        )

        registrations = ExaminationRegistration.objects.filter(
            examination=selected_examination,
            registration_status="APPROVED",
            semester_registration__semester=assignment.subject.semester,
            semester_registration__curriculum=assignment.curriculum,
        ).select_related(
            "student",
            "semester_registration",
        )

        existing_results = ResultEntry.objects.filter(
            examination_registration__in=registrations,
            subject=assignment.subject,
        )

        result_map = {
            result.examination_registration_id: result
            for result in existing_results
        }

        for registration in registrations:

            registration.result_entry = result_map.get(registration.id)

            registration.is_locked = (
                registration.result_entry is not None
                and registration.result_entry.workflow_status in [
                    "SUBMITTED",
                    "APPROVED",
                ]
            )

            registration.class_test_marks = ""
            registration.attendance_assignment_marks = ""
            registration.theory_external_marks = ""
            registration.practical_internal_marks = ""
            registration.practical_external_marks = ""

            registration.total_marks = ""
            registration.grade = ""
            registration.result_status = ""

            registration.current_marks = ""

            if registration.result_entry:
                registration.class_test_marks = (
                    registration.result_entry.class_test_marks
                )

                registration.attendance_assignment_marks = (
                    registration.result_entry.attendance_assignment_marks
                )

                registration.theory_external_marks = (
                    registration.result_entry.theory_external_marks
                )

                registration.practical_internal_marks = (
                    registration.result_entry.practical_internal_marks
                )

                registration.practical_external_marks = (
                    registration.result_entry.practical_external_marks
                )

                registration.total_marks = (
                    registration.result_entry.total_marks
                )

                registration.grade = (
                    registration.result_entry.grade
                )

                registration.result_status = (
                    registration.result_entry.result_status
                )

            if active_window and registration.result_entry:
                if active_window.assessment_type == "CLASS_TEST":

                    registration.current_marks = (
                        registration.result_entry.class_test_marks
                    )

                elif active_window.assessment_type == "ATTENDANCE_ASSIGNMENT":

                    registration.current_marks = (
                        registration.result_entry.attendance_assignment_marks
                    )

                elif active_window.assessment_type == "END_SEMESTER":

                    registration.current_marks = (
                        registration.result_entry.theory_external_marks
                    )

    if request.method == "POST":
        action = request.POST.get("action")
        for registration in registrations:

            if active_window.assessment_type == "CLASS_TEST":

                marks = request.POST.get(
                    f"class_test_{registration.id}"
                )

            elif active_window.assessment_type == "ATTENDANCE_ASSIGNMENT":

                marks = request.POST.get(
                    f"attendance_assignment_{registration.id}"
                )

            elif active_window.assessment_type == "END_SEMESTER":

                theory_external_marks = request.POST.get(
                    f"theory_external_{registration.id}"
                )

                practical_internal_marks = request.POST.get(
                    f"practical_internal_{registration.id}"
                )

                practical_external_marks = request.POST.get(
                    f"practical_external_{registration.id}"
                )

            else:

                marks = None

            if active_window.assessment_type != "END_SEMESTER":

                if marks in [None, ""]:
                    continue

                marks = float(marks)

            else:

                if (
                    theory_external_marks in [None, ""]
                    and practical_internal_marks in [None, ""]
                    and practical_external_marks in [None, ""]
                ):
                    continue

            if active_window.assessment_type == "CLASS_TEST":

                max_marks = assignment.subject.class_test_max

                if marks > max_marks:
                    messages.error(
                        request,
                        f"Class Test marks cannot exceed {max_marks} for {registration.student.roll_number}."
                    )
                    return redirect(
                        f"{request.path}?examination={selected_examination.id}"
                    )

            elif active_window.assessment_type == "ATTENDANCE_ASSIGNMENT":

                max_marks = assignment.subject.attendance_assignment_max

                if marks > max_marks:
                    messages.error(
                        request,
                        f"Attendance & Assignment marks cannot exceed {max_marks} for {registration.student.roll_number}."
                    )
                    return redirect(
                        f"{request.path}?examination={selected_examination.id}"
                    )

            elif active_window.assessment_type == "END_SEMESTER":

                if theory_external_marks not in [None, ""]:

                    theory_external_marks = float(theory_external_marks)

                    if theory_external_marks > assignment.subject.theory_external_max:
                        messages.error(
                            request,
                            f"Theory External marks cannot exceed {assignment.subject.theory_external_max} for {registration.student.roll_number}."
                        )
                        return redirect(
                            f"{request.path}?examination={selected_examination.id}"
                        )

                if practical_internal_marks not in [None, ""]:

                    practical_internal_marks = float(practical_internal_marks)

                    if practical_internal_marks > assignment.subject.practical_internal_max:
                        messages.error(
                            request,
                            f"Practical Internal marks cannot exceed {assignment.subject.practical_internal_max} for {registration.student.roll_number}."
                        )
                        return redirect(
                            f"{request.path}?examination={selected_examination.id}"
                        )

                if practical_external_marks not in [None, ""]:

                    practical_external_marks = float(practical_external_marks)

                    if practical_external_marks > assignment.subject.practical_external_max:
                        messages.error(
                            request,
                            f"Practical External marks cannot exceed {assignment.subject.practical_external_max} for {registration.student.roll_number}."
                        )
                        return redirect(
                            f"{request.path}?examination={selected_examination.id}"
                        )

            result_entry, created = ResultEntry.objects.get_or_create(
                examination_registration=registration,
                subject=assignment.subject,
                defaults={
                   "faculty_assignment": assignment,
                },
            )
            print(
                "DEBUG:",
                registration.student.roll_number,
                result_entry.id,
                result_entry.workflow_status,
            )
            if result_entry.faculty_assignment is None:
                result_entry.faculty_assignment = assignment

            if result_entry.workflow_status in ["SUBMITTED", "APPROVED"]:
                continue

            if active_window.assessment_type == "CLASS_TEST":

                result_entry.class_test_marks = marks

            elif active_window.assessment_type == "ATTENDANCE_ASSIGNMENT":

                result_entry.attendance_assignment_marks = marks

            elif active_window.assessment_type == "END_SEMESTER":

                if theory_external_marks not in [None, ""]:
                    result_entry.theory_external_marks = theory_external_marks

                if practical_internal_marks not in [None, ""]:
                    result_entry.practical_internal_marks = practical_internal_marks

                if practical_external_marks not in [None, ""]:
                    result_entry.practical_external_marks = practical_external_marks

            result_entry.save()

        if action == "submit":

            for registration in registrations:

                ResultEntry.objects.filter(
                    examination_registration=registration,
                    subject=assignment.subject,
                    workflow_status__in=["DRAFT", "RETURNED"],
                ).update(
                    workflow_status="SUBMITTED"
                )

            messages.success(
                request,
                "Marks submitted successfully."
            )

        else:

            messages.success(
                request,
                "Marks saved successfully."
            )

        return redirect(
            f"{request.path}?examination={selected_examination.id}"
        )
    returned_entry = ResultEntry.objects.filter(
        faculty_assignment=assignment,
        examination_registration__examination=selected_examination,
        workflow_status="RETURNED",
    ).first()
    return render(
        request,
        "accounts/enter_marks.html",
        {
            "assignment": assignment,
            "registrations": registrations,
            "examinations": examinations,
            "selected_examination": selected_examination,
            "active_window": active_window,
            "returned_entry": returned_entry,
        },
    )

@login_required
def backlog_enter_marks(request, assignment_id):

    assignment = get_object_or_404(
        BacklogFacultyAssignment,
        id=assignment_id,
        faculty__user=request.user,
        is_active=True,
    )

    # Examination comes directly from the faculty assignment.
    # Faculty should NOT select another examination.
    selected_examination = assignment.examination

    subject = assignment.subject
    subject_type = subject.subject_type

    # ---------------------------------------------------------
    # Determine the currently active assessment window
    # ---------------------------------------------------------

    active_window = (
        AssessmentWindow.objects.filter(
            examination=selected_examination,
            start_date__lte=timezone.localdate(),
            end_date__gte=timezone.localdate(),
        )
        .order_by("start_date", "id")
        .first()
    )

    # ---------------------------------------------------------
    # Approved backlog students for this subject
    # ---------------------------------------------------------

    registrations = (
        BacklogRegistrationSubject.objects.filter(
            backlog_registration__examination=selected_examination,
            backlog_registration__registration_status="APPROVED",
            subject=subject,
        )
        .select_related(
            "backlog_registration__student",
            "subject",
        )
    )

    # ---------------------------------------------------------
    # Existing backlog result entries
    # ---------------------------------------------------------

    existing_results = BacklogResultEntry.objects.filter(
        backlog_registration_subject__in=registrations,
        subject=subject,
    )

    result_map = {
        result.backlog_registration_subject_id: result
        for result in existing_results
    }

    for registration_subject in registrations:

        registration_subject.result_entry = result_map.get(
            registration_subject.id
        )

        registration_subject.is_locked = (
            registration_subject.result_entry is not None
            and registration_subject.result_entry.workflow_status
            in [
                "SUBMITTED",
                "APPROVED",
            ]
        )

        registration_subject.class_test_marks = ""
        registration_subject.attendance_assignment_marks = ""
        registration_subject.theory_external_marks = ""
        registration_subject.practical_internal_marks = ""
        registration_subject.practical_external_marks = ""

        registration_subject.total_marks = ""
        registration_subject.percentage = ""
        registration_subject.grade = ""
        registration_subject.result_status = ""

        if registration_subject.result_entry:

            entry = registration_subject.result_entry

            registration_subject.class_test_marks = (
                entry.class_test_marks
            )

            registration_subject.attendance_assignment_marks = (
                entry.attendance_assignment_marks
            )

            registration_subject.theory_external_marks = (
                entry.theory_external_marks
            )

            registration_subject.practical_internal_marks = (
                entry.practical_internal_marks
            )

            registration_subject.practical_external_marks = (
                entry.practical_external_marks
            )

            registration_subject.total_marks = (
                entry.total_marks
            )

            registration_subject.percentage = (
                entry.percentage
            )

            registration_subject.grade = (
                entry.grade
            )

            registration_subject.result_status = (
                entry.result_status
            )

    # ---------------------------------------------------------
    # Determine which component is currently editable
    # ---------------------------------------------------------

    editable_class_test = False
    editable_attendance_assignment = False
    editable_end_semester = False

    if active_window:

        if active_window.assessment_type == "CLASS_TEST":

            if subject_type in [
                "THEORY",
                "THEORY_PRACTICAL",
            ]:
                editable_class_test = True

        elif (
            active_window.assessment_type
            == "ATTENDANCE_ASSIGNMENT"
        ):

            if subject_type in [
                "THEORY",
                "THEORY_PRACTICAL",
            ]:
                editable_attendance_assignment = True

        elif active_window.assessment_type == "END_SEMESTER":

            if subject_type in [
                "THEORY",
                "PRACTICAL",
                "THEORY_PRACTICAL",
            ]:
                editable_end_semester = True

    # ---------------------------------------------------------
    # POST
    # ---------------------------------------------------------

    if request.method == "POST":

        action = request.POST.get("action")

        # No active window = nothing can be edited.
        if not active_window:

            messages.error(
                request,
                "There is no active assessment window. "
                "Marks cannot be entered or modified.",
            )

            return redirect(request.path)

        for registration_subject in registrations:

            entry, created = (
                BacklogResultEntry.objects.get_or_create(
                    backlog_registration_subject=registration_subject,
                    defaults={
                        "subject": subject,
                        "faculty_assignment": assignment,
                    },
                )
            )

            # Submitted/approved records are permanently locked
            # for faculty editing.
            if entry.workflow_status in [
                "SUBMITTED",
                "APPROVED",
            ]:
                continue

            student_roll = (
                registration_subject
                .backlog_registration
                .student
                .roll_number
            )

            # =================================================
            # CLASS TEST
            # =================================================

            if editable_class_test:

                marks = request.POST.get(
                    f"class_test_{registration_subject.id}"
                )

                if marks not in [None, ""]:

                    try:
                        marks = float(marks)
                    except ValueError:

                        messages.error(
                            request,
                            f"Invalid Class Test marks "
                            f"for {student_roll}.",
                        )

                        return redirect(request.path)

                    if marks < 0:

                        messages.error(
                            request,
                            f"Class Test marks cannot be negative "
                            f"for {student_roll}.",
                        )

                        return redirect(request.path)

                    if marks > subject.class_test_max:

                        messages.error(
                            request,
                            f"Class Test marks cannot exceed "
                            f"{subject.class_test_max} "
                            f"for {student_roll}.",
                        )

                        return redirect(request.path)

                    entry.class_test_marks = marks

            # =================================================
            # ATTENDANCE & ASSIGNMENT
            # =================================================

            elif editable_attendance_assignment:

                marks = request.POST.get(
                    f"attendance_assignment_{registration_subject.id}"
                )

                if marks not in [None, ""]:

                    try:
                        marks = float(marks)
                    except ValueError:

                        messages.error(
                            request,
                            f"Invalid Attendance & Assignment marks "
                            f"for {student_roll}.",
                        )

                        return redirect(request.path)

                    if marks < 0:

                        messages.error(
                            request,
                            f"Attendance & Assignment marks cannot "
                            f"be negative for {student_roll}.",
                        )

                        return redirect(request.path)

                    if (
                        marks
                        > subject.attendance_assignment_max
                    ):

                        messages.error(
                            request,
                            f"Attendance & Assignment marks cannot "
                            f"exceed "
                            f"{subject.attendance_assignment_max} "
                            f"for {student_roll}.",
                        )

                        return redirect(request.path)

                    entry.attendance_assignment_marks = marks

            # =================================================
            # END SEMESTER
            # =================================================

            elif editable_end_semester:

                # ---------------------------------------------
                # Theory External
                # ---------------------------------------------

                if subject_type in [
                    "THEORY",
                    "THEORY_PRACTICAL",
                ]:

                    marks = request.POST.get(
                        f"theory_external_{registration_subject.id}"
                    )

                    if marks not in [None, ""]:

                        try:
                            marks = float(marks)
                        except ValueError:

                            messages.error(
                                request,
                                f"Invalid Theory External marks "
                                f"for {student_roll}.",
                            )

                            return redirect(request.path)

                        if marks < 0:

                            messages.error(
                                request,
                                f"Theory External marks cannot be "
                                f"negative for {student_roll}.",
                            )

                            return redirect(request.path)

                        if marks > subject.theory_external_max:

                            messages.error(
                                request,
                                f"Theory External marks cannot exceed "
                                f"{subject.theory_external_max} "
                                f"for {student_roll}.",
                            )

                            return redirect(request.path)

                        entry.theory_external_marks = marks

                # ---------------------------------------------
                # Practical Internal
                # ---------------------------------------------

                if subject_type in [
                    "PRACTICAL",
                    "THEORY_PRACTICAL",
                ]:

                    marks = request.POST.get(
                        f"practical_internal_{registration_subject.id}"
                    )

                    if marks not in [None, ""]:

                        try:
                            marks = float(marks)
                        except ValueError:

                            messages.error(
                                request,
                                f"Invalid Practical Internal marks "
                                f"for {student_roll}.",
                            )

                            return redirect(request.path)

                        if marks < 0:

                            messages.error(
                                request,
                                f"Practical Internal marks cannot be "
                                f"negative for {student_roll}.",
                            )

                            return redirect(request.path)

                        if (
                            marks
                            > subject.practical_internal_max
                        ):

                            messages.error(
                                request,
                                f"Practical Internal marks cannot exceed "
                                f"{subject.practical_internal_max} "
                                f"for {student_roll}.",
                            )

                            return redirect(request.path)

                        entry.practical_internal_marks = marks

                    # -----------------------------------------
                    # Practical External
                    # -----------------------------------------

                    marks = request.POST.get(
                        f"practical_external_{registration_subject.id}"
                    )

                    if marks not in [None, ""]:

                        try:
                            marks = float(marks)
                        except ValueError:

                            messages.error(
                                request,
                                f"Invalid Practical External marks "
                                f"for {student_roll}.",
                            )

                            return redirect(request.path)

                        if marks < 0:

                            messages.error(
                                request,
                                f"Practical External marks cannot be "
                                f"negative for {student_roll}.",
                            )

                            return redirect(request.path)

                        if (
                            marks
                            > subject.practical_external_max
                        ):

                            messages.error(
                                request,
                                f"Practical External marks cannot exceed "
                                f"{subject.practical_external_max} "
                                f"for {student_roll}.",
                            )

                            return redirect(request.path)

                        entry.practical_external_marks = marks

            entry.save()

        # -----------------------------------------------------
        # SUBMIT
        # -----------------------------------------------------

        if action == "submit":

            BacklogResultEntry.objects.filter(
                backlog_registration_subject__in=registrations,
                subject=subject,
                workflow_status__in=[
                    "DRAFT",
                    "RETURNED",
                ],
            ).update(
                workflow_status="SUBMITTED",
                submitted_at=timezone.now(),
            )

            messages.success(
                request,
                "Backlog marks submitted successfully.",
            )

        else:

            messages.success(
                request,
                "Backlog marks saved successfully.",
            )

        return redirect(request.path)

    return render(
        request,
        "accounts/backlog_enter_marks.html",
        {
            "assignment": assignment,
            "registrations": registrations,
            "selected_examination": selected_examination,
            "active_window": active_window,
            "editable_class_test": editable_class_test,
            "editable_attendance_assignment":
                editable_attendance_assignment,
            "editable_end_semester": editable_end_semester,
        },
    )

@login_required
def coe_marks_verification(request):

    submission_summary = []

    # =========================================================
    # REGULAR FACULTY SUBMISSIONS
    # =========================================================

    regular_entries = (
        ResultEntry.objects.filter(
            workflow_status="SUBMITTED",
        )
        .select_related(
            "faculty_assignment__faculty__user",
            "faculty_assignment__subject",
            "faculty_assignment__curriculum",
            "faculty_assignment__academic_year",
            "examination_registration__examination",
        )
    )

    regular_map = {}

    for entry in regular_entries:

        key = (
            entry.faculty_assignment.id,
            entry.examination_registration.examination.id,
        )

        if key not in regular_map:

            regular_map[key] = {
                "submission_type": "REGULAR",

                "faculty_assignment_id":
                    entry.faculty_assignment.id,

                "examination_id":
                    entry.examination_registration.examination.id,

                "faculty":
                    entry.faculty_assignment.faculty,

                "subject":
                    entry.faculty_assignment.subject,

                "academic_year":
                    entry.faculty_assignment.academic_year,

                "curriculum":
                    entry.faculty_assignment.curriculum,

                "examination":
                    entry.examination_registration.examination,

                "student_count": 0,

                "status":
                    entry.workflow_status,
            }

        regular_map[key]["student_count"] += 1

    submission_summary.extend(
        regular_map.values()
    )

    # =========================================================
    # BACKLOG FACULTY SUBMISSIONS
    # =========================================================

    backlog_entries = (
        BacklogResultEntry.objects.filter(
            workflow_status="SUBMITTED",
        )
        .select_related(
            "faculty_assignment__faculty__user",
            "faculty_assignment__curriculum",
            "faculty_assignment__subject",
            "faculty_assignment__examination",
            "backlog_registration_subject__backlog_registration__student",
        )
    )

    backlog_map = {}

    for entry in backlog_entries:

        key = (
            entry.faculty_assignment.id,
            entry.faculty_assignment.examination.id,
        )

        if key not in backlog_map:

            backlog_map[key] = {
                "submission_type": "BACKLOG",

                "faculty_assignment_id":
                    entry.faculty_assignment.id,

                "examination_id":
                    entry.faculty_assignment.examination.id,

                "faculty":
                    entry.faculty_assignment.faculty,

                "subject":
                    entry.faculty_assignment.subject,

                "academic_year": None,

                "curriculum":
                    entry.faculty_assignment.curriculum,

                "examination":
                    entry.faculty_assignment.examination,

                "student_count": 0,

                "status":
                    entry.workflow_status,
            }

        backlog_map[key]["student_count"] += 1

    submission_summary.extend(
        backlog_map.values()
    )

    return render(
        request,
        "accounts/coe_marks_verification.html",
        {
            "submitted_assignments":
                submission_summary,
        },
    )

@login_required
def coe_verify_submission(
    request,
    faculty_assignment_id,
    examination_id,
):

    entries = ResultEntry.objects.filter(
        faculty_assignment_id=faculty_assignment_id,
        examination_registration__examination_id=examination_id,
    )
    entry = entries.first()

    if not entry:
        return HttpResponse("No submitted records found.")

    return render(
        request,
        "accounts/coe_verify_submission.html",
        {
            "entry": entry,
            "entries": entries,
            "faculty_assignment_id": faculty_assignment_id,
            "examination_id": examination_id,
        },
    )

@login_required
def coe_approve_submission(
    request,
    faculty_assignment_id,
    examination_id,
):

    entries = ResultEntry.objects.filter(
        faculty_assignment_id=faculty_assignment_id,
        examination_registration__examination_id=examination_id,
    )

    entry = entries.first()

    if not entry:
        return HttpResponse("Submission not found.")

    if request.method == "POST":

        entries.update(
            workflow_status="APPROVED",
            approved_at=timezone.now(),
        )

        return redirect("coe_marks_verification")

    return render(
        request,
        "accounts/coe_approve_submission.html",
        {
            "entry": entry,
            "entries": entries,
            "faculty_assignment_id": faculty_assignment_id,
            "examination_id": examination_id,
        },
    )

@login_required
def coe_return_submission(
    request,
    faculty_assignment_id,
    examination_id,
):

    entries = ResultEntry.objects.filter(
        faculty_assignment_id=faculty_assignment_id,
        examination_registration__examination_id=examination_id,
    )

    entry = entries.first()

    if request.method == "POST":

        remarks = request.POST.get("remarks")

        entries.update(
            workflow_status="RETURNED",
            workflow_remarks=remarks,
        )

        return redirect("coe_marks_verification")

    if not entry:
        return HttpResponse("Submission not found.")

    return render(
        request,
        "accounts/coe_return_submission.html",
        {
            "entry": entry,
            "entries": entries,
            "faculty_assignment_id": faculty_assignment_id,
            "examination_id": examination_id,
        },
    )

@login_required
def coe_verify_backlog_submission(
    request,
    faculty_assignment_id,
    examination_id,
):

    entries = BacklogResultEntry.objects.filter(
        faculty_assignment_id=faculty_assignment_id,
        backlog_registration_subject__backlog_registration__examination_id=examination_id,
    ).select_related(
        "backlog_registration_subject__backlog_registration__student",
        "subject",
        "faculty_assignment__faculty",
        "faculty_assignment__curriculum",
        "faculty_assignment__examination",
    )

    entry = entries.first()

    if not entry:
        return HttpResponse(
            "No submitted backlog records found."
        )

    return render(
        request,
        "accounts/coe_verify_backlog_submission.html",
        {
            "entry": entry,
            "entries": entries,
            "faculty_assignment_id": faculty_assignment_id,
            "examination_id": examination_id,
        },
    )

@login_required
def coe_approve_backlog_submission(
    request,
    faculty_assignment_id,
    examination_id,
):

    entries = BacklogResultEntry.objects.filter(
        faculty_assignment_id=faculty_assignment_id,
        backlog_registration_subject__backlog_registration__examination_id=examination_id,
    )

    entry = entries.first()

    if not entry:
        return HttpResponse(
            "Backlog submission not found."
        )

    if request.method == "POST":

        entries.update(
            workflow_status="APPROVED",
            approved_at=timezone.now(),
        )
        
        return redirect(
            "coe_marks_verification"
        )

    return render(
        request,
        "accounts/coe_approve_backlog_submission.html",
        {
            "entry": entry,
            "entries": entries,
            "faculty_assignment_id": faculty_assignment_id,
            "examination_id": examination_id,
        },
    )

@login_required
def coe_return_backlog_submission(
    request,
    faculty_assignment_id,
    examination_id,
):

    entries = BacklogResultEntry.objects.filter(
        faculty_assignment_id=faculty_assignment_id,
        backlog_registration_subject__backlog_registration__examination_id=examination_id,
    )

    entry = entries.first()

    if request.method == "POST":

        remarks = request.POST.get(
            "remarks",
            "",
        )

        entries.update(
            workflow_status="RETURNED",
            workflow_remarks=remarks,
        )

        return redirect(
            "coe_marks_verification"
        )

    if not entry:
        return HttpResponse(
            "Backlog submission not found."
        )

    return render(
        request,
        "accounts/coe_return_backlog_submission.html",
        {
            "entry": entry,
            "entries": entries,
            "faculty_assignment_id": faculty_assignment_id,
            "examination_id": examination_id,
        },
    )

@login_required
def result_processing(request):

    approved_assignments = (
        ResultEntry.objects.filter(
            workflow_status="APPROVED",
        )
        .select_related(
            "faculty_assignment__faculty__user",
            "faculty_assignment__subject",
            "faculty_assignment__curriculum",
            "faculty_assignment__academic_year",
            "examination_registration__examination",
        )
    )

    processing_summary = []
    processing_map = {}

    for entry in approved_assignments:

        key = (
            entry.faculty_assignment.id,
            entry.examination_registration.examination.id,
        )

        if key not in processing_map:

            processing_map[key] = {
                "faculty_assignment_id": entry.faculty_assignment.id,
                "examination_id": entry.examination_registration.examination.id,

                "faculty": entry.faculty_assignment.faculty,
                "subject": entry.faculty_assignment.subject,
                "academic_year": entry.faculty_assignment.academic_year,
                "curriculum": entry.faculty_assignment.curriculum,
                "examination": entry.examination_registration.examination,

                "student_count": 0,
                "status": entry.workflow_status,
            }

        processing_map[key]["student_count"] += 1

    processing_summary = list(processing_map.values())

    return render(
        request,
        "accounts/result_processing.html",
        {
            "processing_summary": processing_summary,
        },
    )

@login_required
def backlog_result_processing(request):

    approved_assignments = (
        BacklogResultEntry.objects.filter(
            workflow_status="APPROVED",
        )
        .select_related(
            "faculty_assignment__faculty__user",
            "faculty_assignment__subject",
            "faculty_assignment__curriculum",
            "faculty_assignment__examination",
        )
    )

    processing_summary = []

    processing_map = {}

    for entry in approved_assignments:

        key = (
            entry.faculty_assignment.id,
            entry.faculty_assignment.examination.id,
        )

        if key not in processing_map:

            processing_map[key] = {
                "faculty_assignment_id":
                    entry.faculty_assignment.id,

                "examination_id":
                    entry.faculty_assignment.examination.id,

                "faculty":
                    entry.faculty_assignment.faculty,

                "subject":
                    entry.faculty_assignment.subject,

                "curriculum":
                    entry.faculty_assignment.curriculum,

                "examination":
                    entry.faculty_assignment.examination,

                "student_count":
                    0,

                "status":
                    entry.workflow_status,
            }

        processing_map[key]["student_count"] += 1

    processing_summary = list(
        processing_map.values()
    )

    return render(
        request,
        "accounts/backlog_result_processing.html",
        {
            "processing_summary":
                processing_summary,
        },
    )

@login_required
def process_backlog_result(
    request,
    faculty_assignment_id,
    examination_id,
):

    entries = BacklogResultEntry.objects.filter(
        faculty_assignment_id=faculty_assignment_id,
        backlog_registration_subject__backlog_registration__examination_id=examination_id,
    ).select_related(
        "subject",
        "backlog_registration_subject__backlog_registration__student",
        "backlog_registration_subject__backlog_registration__examination",
    )

    entry = entries.first()

    if not entry:
        return HttpResponse(
            "Backlog Result not found."
        )

    if request.method == "POST":

        for entry in entries:

            calculate_backlog_result(entry)

            entry.save(
                update_fields=[
                    "total_marks",
                    "percentage",
                    "grade",
                    "grade_point",
                    "result_status",
                ]
            )

            student = (
                entry.backlog_registration_subject
                .backlog_registration
                .student
            )

            examination = (
                entry.backlog_registration_subject
                .backlog_registration
                .examination
            )

            attempt_number = get_live_attempt_number(
                student=student,
                subject=entry.subject,
                examination=examination,
            )

            StudentAcademicResult.objects.update_or_create(

                student=student,

                examination=examination,

                subject=entry.subject,

                defaults={

                    "attempt_number":
                        attempt_number,

                    "class_test_marks":
                        entry.class_test_marks,

                    "attendance_assignment_marks":
                        entry.attendance_assignment_marks,

                    "theory_external_marks":
                        entry.theory_external_marks,

                    "practical_internal_marks":
                        entry.practical_internal_marks,

                    "practical_external_marks":
                        entry.practical_external_marks,

                    "total_marks":
                        entry.total_marks,

                    "percentage":
                        entry.percentage,

                    "grade":
                        entry.grade,

                    "grade_point":
                        entry.grade_point,

                    "result_status":
                        entry.result_status,

                    "source":
                        "LIVE",
                },
            )

        messages.success(
            request,
            "Backlog result processed successfully."
        )

        return redirect(
            "backlog_result_processing"
        )

    return render(
        request,
        "accounts/process_backlog_result.html",
        {
            "entry": entry,
            "entries": entries,
            "faculty_assignment_id":
                faculty_assignment_id,
            "examination_id":
                examination_id,
        },
    )

@login_required
def process_result(
    request,
    faculty_assignment_id,
    examination_id,
):

    entries = ResultEntry.objects.filter(
        faculty_assignment_id=faculty_assignment_id,
        examination_registration__examination_id=examination_id,
    ).select_related("subject")

    entry = entries.first()

    if not entry:
        return HttpResponse("Result not found.")

    if request.method == "POST":

        for entry in entries:

            subject = entry.subject

            # Calculation code will come here
            theory_internal = (
                (entry.class_test_marks or 0)
                + (entry.attendance_assignment_marks or 0)
            )

            total_marks = (
                theory_internal
                + (entry.theory_external_marks or 0)
                + (entry.practical_internal_marks or 0)
                + (entry.practical_external_marks or 0)
            )

            maximum_marks = (
                subject.theory_internal_max
                + subject.theory_external_max
                + subject.practical_internal_max
                + subject.practical_external_max
            )

            if maximum_marks > 0:
                percentage = (
                    total_marks * 100
                ) / maximum_marks
            else:
                percentage = 0

            entry.total_marks = total_marks
            entry.percentage = percentage
                        # Default values
            entry.grade = ""
            entry.grade_point = None
            entry.result_status = ""

            # Rule 1: Overall percentage must be at least 40
            if percentage < 40:

                entry.grade = "FF"
                entry.grade_point = 0
                entry.result_status = "FAIL"

            else:

                component_failed = False

                # Theory subjects
                if subject.subject_type in ["THEORY", "THEORY_PRACTICAL"]:

                    if theory_internal < subject.theory_internal_min:
                        component_failed = True

                    if (
                        (entry.theory_external_marks or 0)
                        < subject.theory_external_min
                    ):
                        component_failed = True

                # Practical subjects
                if subject.subject_type in ["PRACTICAL", "THEORY_PRACTICAL"]:

                    if (
                        (entry.practical_internal_marks or 0)
                        < subject.practical_internal_min
                    ):
                        component_failed = True

                    if (
                        (entry.practical_external_marks or 0)
                        < subject.practical_external_min
                    ):
                        component_failed = True

                if component_failed:

                    entry.grade = "FF"
                    entry.grade_point = 0
                    entry.result_status = "FAIL"

                else:

                    if percentage >= 85:
                        entry.grade = "A+"
                        entry.grade_point = 10

                    elif percentage >= 75:
                        entry.grade = "A"
                        entry.grade_point = 9

                    elif percentage >= 65:
                        entry.grade = "B+"
                        entry.grade_point = 8

                    elif percentage >= 55:
                        entry.grade = "B"
                        entry.grade_point = 7

                    elif percentage >= 45:
                        entry.grade = "C"
                        entry.grade_point = 6

                    else:
                        entry.grade = "D"
                        entry.grade_point = 5

                    entry.result_status = "PASS"

            entry.save(
                update_fields=[
                    "total_marks",
                    "percentage",
                    "grade",
                    "grade_point",
                    "result_status",
                ]
            )

            attempt_number = get_live_attempt_number(
                student=entry.examination_registration.student,
                subject=entry.subject,
                examination=entry.examination_registration.examination,
            )

            StudentAcademicResult.objects.update_or_create(

                student=entry.examination_registration.student,

                examination=entry.examination_registration.examination,

                subject=entry.subject,

                defaults={

                    "attempt_number":
                        attempt_number,

                    "class_test_marks":
                        entry.class_test_marks,

                    "attendance_assignment_marks":
                        entry.attendance_assignment_marks,

                    "theory_external_marks":
                        entry.theory_external_marks,

                    "practical_internal_marks":
                        entry.practical_internal_marks,

                    "practical_external_marks":
                        entry.practical_external_marks,

                    "total_marks":
                        entry.total_marks,

                    "percentage":
                        entry.percentage,

                    "grade":
                        entry.grade,

                    "grade_point":
                        entry.grade_point,

                    "result_status":
                        entry.result_status,

                    "source":
                        "LIVE",
                },
            )

        messages.success(
            request,
            "Result processed successfully."
        )

        return redirect("result_processing")

    return render(
        request,
        "accounts/process_result.html",
        {
            "entry": entry,
            "entries": entries,
            "faculty_assignment_id": faculty_assignment_id,
            "examination_id": examination_id,
        },
    )

@login_required
def result_publication(request):

    published_map = {}

    entries = (
        ResultEntry.objects.filter(
            workflow_status="APPROVED",
        )
        .select_related(
            "faculty_assignment__faculty",
            "faculty_assignment__subject",
            "faculty_assignment__academic_year",
            "faculty_assignment__curriculum",
            "examination_registration__examination",
        )
    )

    for entry in entries:

        key = (
            entry.faculty_assignment.id,
            entry.examination_registration.examination.id,
        )

        if key not in published_map:

            published_map[key] = {

                "faculty_assignment_id": entry.faculty_assignment.id,

                "examination_id":
                entry.examination_registration.examination.id,

                "faculty":
                entry.faculty_assignment.faculty,

                "subject":
                entry.faculty_assignment.subject,

                "academic_year":
                entry.faculty_assignment.academic_year,

                "curriculum":
                entry.faculty_assignment.curriculum,

                "examination":
                entry.examination_registration.examination,

                "student_count": 0,

                "published": entry.is_published,
            }

        published_map[key]["student_count"] += 1

    return render(
        request,
        "accounts/result_publication.html",
        {
            "publication_summary":
            list(published_map.values())
        },
    )

@login_required
def publish_result(
    request,
    faculty_assignment_id,
    examination_id,
):

    ResultEntry.objects.filter(
        faculty_assignment_id=faculty_assignment_id,
        examination_registration__examination_id=examination_id,
        workflow_status="APPROVED",
        is_published=False,
    ).update(
        is_published=True,
        published_at=timezone.now(),
        published_by=request.user,
    )

    return redirect("result_publication")

@login_required
def semester_result_processing(request):

    examinations = Examination.objects.all().order_by("id")

    students = []

    selected_examination = None

    selected_semester = None

    if request.method == "POST":

        examination = Examination.objects.get(
            id=request.POST.get("examination")
        )

        semester = int(
            request.POST.get("semester")
        )

        students = (
            Student.objects.filter(
                academic_results__examination=examination,
                academic_results__subject__semester=semester,
            )
            .distinct()
            .order_by("roll_number")
        )

        processed = 0

        for student in students:

            calculate_sgpa(
                student,
                examination,
                semester,
            )

            processed += 1

        messages.success(
            request,
            f"{processed} students processed successfully."
        )

        return redirect(
            f"{request.path}?examination={examination.id}&semester={semester}"
        )

    if request.method == "GET":

        examination_id = request.GET.get("examination")

        semester = request.GET.get("semester")

        if examination_id and semester:

            selected_examination = Examination.objects.get(
                id=examination_id
            )

            selected_semester = int(semester)

            students = (
                Student.objects.filter(
                    academic_results__examination=selected_examination,
                    academic_results__subject__semester=selected_semester,
                )
                .distinct()
                .order_by("roll_number")
            )

            semester_result_map = {

                result.student_id: result

                for result in SemesterResult.objects.filter(
                    examination=selected_examination,
                    semester=selected_semester,
                )
            }

            for student in students:

                student.semester_result = (
                    semester_result_map.get(student.id)
                )

    return render(
        request,
        "accounts/semester_result_processing.html",
        {
            "examinations": examinations,
            "students": students,
            "selected_examination": selected_examination,
            "selected_semester": selected_semester,
        },
    )

@login_required
def combined_result_processing(request):

    examinations = Examination.objects.all().order_by("id")

    students = []

    selected_examination = None

    if request.method == "POST":

        examination = Examination.objects.get(
            id=request.POST.get("examination")
        )

        students = (
            Student.objects.filter(
                semester_results__examination=examination,
            )
            .distinct()
            .order_by("roll_number")
        )

        processed = 0

        for student in students:

            calculate_cgpa(
                student,
                examination,
            )

            processed += 1

        messages.success(
            request,
            f"{processed} combined results processed successfully."
        )

        return redirect(
            f"{request.path}?examination={examination.id}"
        )

    if request.method == "GET":

        examination_id = request.GET.get("examination")

        if examination_id:

            selected_examination = Examination.objects.get(
                id=examination_id
            )

            students = (
                Student.objects.filter(
                    semester_results__examination=selected_examination,
                )
                .distinct()
                .order_by("roll_number")
            )

            combined_result_map = {

                result.student_id: result

                for result in CombinedResult.objects.filter(
                    examination=selected_examination
                )
            }

            for student in students:

                student.combined_result = (
                    combined_result_map.get(student.id)
                )

    return render(
        request,
        "accounts/combined_result_processing.html",
        {
            "examinations": examinations,
            "students": students,
            "selected_examination": selected_examination,
        },
    )

@login_required
def historical_result_import(request):

    examinations = Examination.objects.all().order_by("id")

    excel_data = None
    invalid_subjects = []      # <-- Add this
    invalid_roll_numbers = []
    invalid_grades = []

    VALID_GRADES = {
        "A+",
        "A",
        "B+",
        "B",
        "C",
        "D",
        "FF",
    }

    if request.method == "POST":

        action = request.POST.get(
            "action",
            "preview",
        )

        if action == "import":

            excel_data = request.session.get(
                "historical_result_excel"
            )

            examination_id = request.session.get(
                "historical_result_examination"
            )

            if not excel_data or not examination_id:

                messages.error(
                    request,
                    "No validated data found. Please preview the Excel again."
                )

                return redirect(
                    "historical_result_import"
                )

            examination = Examination.objects.get(
                id=examination_id
            )

            header = excel_data[0]

            subject_codes = header[1:]

            GRADE_MAP = {
                "A+": (10, "PASS"),
                "A": (9, "PASS"),
                "B+": (8, "PASS"),
                "B": (7, "PASS"),
                "C": (6, "PASS"),
                "D": (5, "PASS"),
                "FF": (0, "FAIL"),
            }

            for row in excel_data[1:]:

                roll_number = str(row[0]).strip()

                student = Student.objects.get(
                    roll_number=roll_number
                )

                for index, grade in enumerate(row[1:]):

                    subject_code = str(
                        subject_codes[index]
                    ).strip()

                    grade = str(
                        grade
                    ).strip().upper()

                    subject = Subject.objects.get(
                        subject_code=subject_code
                    )

                    

                    if grade == "":
                        continue

                    if grade not in GRADE_MAP:
                        continue

                    grade_point, result_status = GRADE_MAP[grade]

                    StudentAcademicResult.objects.update_or_create(

                        student=student,

                        examination=examination,

                        subject=subject,

                        defaults={

                            "grade": grade,

                            "grade_point": grade_point,

                            "result_status": result_status,

                            "source": "HISTORICAL",

                        },
                    )

            request.session.pop(
                "historical_result_excel",
                None,
            )

            request.session.pop(
                "historical_result_examination",
                None,
            )

            messages.success(
                request,
                "Historical results imported successfully."
            )

            return redirect(
                "historical_result_import"
            )

        excel_file = request.FILES.get("excel_file")

        workbook = load_workbook(excel_file)
        worksheet = workbook.active

        excel_data = []

        for row in worksheet.iter_rows(values_only=True):
            excel_data.append(list(row))

        header = excel_data[0]

        subject_codes = header[1:]

        for code in subject_codes:

            code = str(code).strip()

            if not Subject.objects.filter(
                subject_code=code,
                is_active=True
            ).exists():

                invalid_subjects.append(code)

        # Skip header row
        for row in excel_data[1:]:

            roll_number = str(row[0]).strip()

            if not Student.objects.filter(
                roll_number=roll_number,
                is_active=True
            ).exists():

                invalid_roll_numbers.append(roll_number)

        # Skip header row
        for row_index, row in enumerate(excel_data[1:], start=2):

            for col_index, grade in enumerate(row[1:], start=2):

                grade = str(grade).strip().upper()

                if grade != "" and grade not in VALID_GRADES:

                   invalid_grades.append({
                        "row": row_index,
                        "column": col_index,
                        "grade": grade,
                    })

        if (
            action == "preview"
            and not invalid_subjects
            and not invalid_roll_numbers
            and not invalid_grades
        ):

            request.session["historical_result_excel"] = excel_data

            request.session["historical_result_examination"] = (
                request.POST.get("examination")
            )

    return render(
        request,
        "accounts/historical_result_import.html",
        {
            "examinations": examinations,
            "excel_data": excel_data,
            "invalid_subjects": invalid_subjects,
            "invalid_roll_numbers": invalid_roll_numbers,
            "invalid_grades": invalid_grades,
        },
    )

@login_required
def semester_grade_sheet(
    request,
    student_id,
    examination_id,
    semester,
):

    student = Student.objects.get(
        id=student_id
    )

    examination = Examination.objects.get(
        id=examination_id
    )

    # =========================================================
    # RESULT PUBLICATION CHECK
    # =========================================================

    publication = ResultPublication.objects.filter(
        examination=examination,
        semester=semester,
        is_published=True,
    ).first()

    if publication is None:

        messages.error(
            request,
            "Semester Grade Sheet is not available."
        )

        return redirect(
            "student_results"
        )

    # =========================================================
    # OFFICIAL SEMESTER RESULT
    # =========================================================

    semester_result = SemesterResult.objects.get(
        student=student,
        examination=examination,
        semester=semester,
    )

    # =========================================================
    # SUBJECT RESULTS
    # =========================================================
    #
    # IMPORTANT:
    # Do NOT use get_effective_results() here.
    #
    # We need the result corresponding to the
    # published examination.
    #

    results = StudentAcademicResult.objects.filter(
        student=student,
        examination=examination,
        subject__semester=semester,
    ).select_related(
        "subject",
    ).order_by(
        "subject__subject_code",
    )

    return render(
        request,
        "accounts/semester_grade_sheet.html",
        {
            "student": student,
            "examination": examination,
            "semester": semester,
            "results": results,
            "semester_result": semester_result,
        },
    )

@login_required
def combined_grade_sheet(
    request,
    student_id,
    examination_id,
):

    student = Student.objects.get(
        id=student_id
    )

    examination = Examination.objects.get(
        id=examination_id
    )

    combined_result = CombinedResult.objects.get(
        student=student,
        examination=examination,
    )

    # =========================================================
    # PUBLISHED SUBJECT RESULTS
    # =========================================================

    published_examinations = {}

    publications = (
        ResultPublication.objects.filter(
            examination__semester_results__student=student,
            is_published=True,
        )
        .select_related(
            "examination",
        )
        .order_by(
            "semester",
            "-published_on",
            "-id",
        )
    )

    for publication in publications:

        semester = publication.semester

        if semester not in published_examinations:

            published_examinations[semester] = (
                publication.examination
            )

    subject_results = []

    for semester, published_examination in (
        published_examinations.items()
    ):

        semester_results = (
            StudentAcademicResult.objects.filter(
                student=student,
                examination=published_examination,
                subject__semester=semester,
            )
            .select_related(
                "subject",
            )
            .order_by(
                "subject__subject_code",
            )
        )

        subject_results.extend(
            semester_results
        )

    subject_results = sorted(
        subject_results,
        key=lambda result: (
            result.subject.semester,
            result.subject.subject_code,
        ),
    )

    # =========================================================
# CALCULATE DISPLAYED CGPA FROM PUBLISHED RESULTS
# =========================================================

    displayed_registered_credits = Decimal("0")

    displayed_earned_credits = Decimal("0")

    displayed_weighted_grade_points = Decimal("0")

    displayed_result_status = "PASS"

    for result in subject_results:

        credits = result.subject.credits

        displayed_registered_credits += credits

        displayed_weighted_grade_points += (
            credits * result.grade_point
        )

        if result.result_status == "PASS":

            displayed_earned_credits += credits

        else:

            displayed_result_status = "FAIL"

    if displayed_registered_credits > 0:

        displayed_cgpa = (
            displayed_weighted_grade_points
            / displayed_registered_credits
        ).quantize(
            Decimal("0.01")
        )

    else:

        displayed_cgpa = Decimal("0.00")

    # =========================================================
    # BUILD SEMESTER-WISE SUMMARY
    # =========================================================

    semester_data = {}

    for result in subject_results:

        semester = result.subject.semester
        credits = result.subject.credits

        if semester not in semester_data:

            semester_data[semester] = {
                "semester": semester,
                "registered_credits": 0,
                "earned_credits": 0,
                "weighted_grade_points": 0,
                "result_status": "PASS",
            }

        semester_data[semester]["registered_credits"] += (
            credits
        )

        semester_data[semester]["weighted_grade_points"] += (
            credits * result.grade_point
        )

        if result.result_status == "PASS":

            semester_data[semester]["earned_credits"] += (
                credits
            )

        else:

            semester_data[semester]["result_status"] = "FAIL"

    semester_results = []

    for semester in sorted(semester_data):

        data = semester_data[semester]

        registered_credits = data[
            "registered_credits"
        ]

        if registered_credits > 0:

            sgpa = (
                data["weighted_grade_points"]
                / registered_credits
            ).quantize(
                Decimal("0.01")
            )

        else:

            sgpa = Decimal("0.00")

        semester_results.append(
            {
                "semester": semester,
                "registered_credits":
                    registered_credits,
                "earned_credits":
                    data["earned_credits"],
                "sgpa": sgpa,
                "result_status":
                    data["result_status"],
            }
        )

    return render(
        request,
        "accounts/combined_grade_sheet.html",
        {
            "student": student,
            "examination": examination,
            "semester_results": semester_results,
            "combined_result": combined_result,
            "subject_results": subject_results,
            "displayed_cgpa": displayed_cgpa,
            "displayed_result_status": displayed_result_status,

        },
    )

@login_required
def result_publication(request):

    examinations = Examination.objects.all().order_by("id")

    students = []

    publication = None

    selected_examination = None

    selected_semester = None

    if request.method == "POST":

        selected_examination = Examination.objects.get(
            id=request.POST.get("examination")
        )

        selected_semester = int(
            request.POST.get("semester")
        )

        semester_results_exist = SemesterResult.objects.filter(
            examination=selected_examination,
            semester=selected_semester,
        ).exists()

        if not semester_results_exist:

            messages.error(
                request,
                "Semester Result Processing has not been completed for the selected semester."
            )

            return redirect(
                f"{request.path}?examination={selected_examination.id}&semester={selected_semester}"
            )

        combined_results_exist = CombinedResult.objects.filter(
            examination=selected_examination,
        ).exists()

        if not combined_results_exist:

            messages.error(
                request,
                "Combined Result Processing has not been completed."
            )

            return redirect(
                f"{request.path}?examination={selected_examination.id}&semester={selected_semester}"
            )

        publication, created = ResultPublication.objects.get_or_create(
            examination=selected_examination,
            semester=selected_semester,
        )

# Validation 3 - Prevent duplicate publication
        if publication.is_published:

            messages.warning(
                request,
                "Result has already been published."
            )

            return redirect(
                f"{request.path}?examination={selected_examination.id}&semester={selected_semester}"
            )

        publication.is_published = True

        publication.published_by = request.user

        publication.published_on = timezone.now()

        publication.save()

        messages.success(
            request,
            "Result published successfully."
        )

        return redirect(
            f"{request.path}?examination={selected_examination.id}&semester={selected_semester}"
        )

    examination_id = request.GET.get("examination")

    semester = request.GET.get("semester")

    if examination_id and semester:

        selected_examination = Examination.objects.get(
            id=examination_id
        )

        selected_semester = int(semester)

        publication = ResultPublication.objects.filter(
            examination=selected_examination,
            semester=selected_semester,
        ).first()

    publication_records = (
        ResultPublication.objects
        .select_related(
            "examination",
            "published_by",
        )
        .order_by(
            "-examination__id",
            "semester",
        )
    )

    return render(
        request,
        "accounts/result_publication.html",
        {
            "examinations": examinations,
            "publication": publication,
            "publication_records": publication_records,
            "selected_examination": selected_examination,
            "selected_semester": selected_semester,
        },
    )

@login_required
def student_results(request):

    student = Student.objects.get(
        user=request.user
    )

    publications = (
        ResultPublication.objects.filter(
            examination__semester_results__student=student,
            is_published=True,
        )
        .select_related(
            "examination",
        )
        .order_by(
            "semester",
            "-published_on",
            "-id",
        )
    )

    published_results = []

    published_semesters = set()

    for publication in publications:

        semester = publication.semester

        # Only the latest published examination
        # for each semester is displayed.
        if semester in published_semesters:
            continue

        semester_result = SemesterResult.objects.filter(
            student=student,
            examination=publication.examination,
            semester=semester,
        ).first()

        if semester_result is None:
            continue

        published_semesters.add(semester)

        published_results.append(
            {
                "semester_result": semester_result,
                "publication": publication,
            }
        )

    latest_examination = None

    if published_results:

        latest_examination = (
            published_results[-1]["semester_result"].examination
        )

    return render(
        request,
        "accounts/student_results.html",
        {
            "student": student,
            "published_results": published_results,
            "latest_examination": latest_examination,
        },
    )

@login_required
def semester_grade_sheet_list(request):

    students = (
        Student.objects.filter(
            semester_results__isnull=False
        )
        .distinct()
        .order_by("roll_number")
    )

    return render(
        request,
        "accounts/semester_grade_sheet_list.html",
        {
            "students": students,
        },
    )

@login_required
def combined_grade_sheet_list(request):

    students = (
        Student.objects.filter(
            combined_results__isnull=False
        )
        .distinct()
        .order_by("roll_number")
    )

    student_results = []

    for student in students:

        latest_result = (
            CombinedResult.objects.filter(
                student=student
            )
            .order_by("-examination__id")
            .first()
        )

        if not latest_result:
            continue

        # Calculate CGPA from the effective/latest
        # attempt for every subject.
        effective_results = get_effective_results(
            student=student
        )

        registered_credits = Decimal("0")
        weighted_grade_points = Decimal("0")
        result_status = "PASS"

        for result in effective_results:

            credits = result.subject.credits

            registered_credits += credits

            weighted_grade_points += (
                credits * result.grade_point
            )

            if result.result_status == "FAIL":
                result_status = "FAIL"

        if registered_credits > 0:

            displayed_cgpa = (
                weighted_grade_points
                / registered_credits
            ).quantize(
                Decimal("0.01")
            )

        else:

            displayed_cgpa = Decimal("0.00")

        student_results.append(
            {
                "student": student,
                "combined_result": latest_result,
                "displayed_cgpa": displayed_cgpa,
                "displayed_result_status": result_status,
            }
        )

    return render(
        request,
        "accounts/combined_grade_sheet_list.html",
        {
            "student_results": student_results,
        },
    )

@login_required
def student_semester_grade_sheet(
    request,
    examination_id,
    semester,
):

    student = Student.objects.get(
        user=request.user
    )

    publication = (
        ResultPublication.objects.filter(
            examination__semester_results__student=student,
            semester=semester,
            is_published=True,
        )
        .order_by(
            "-published_on",
            "-id",
        )
        .first()
    )

    if publication is None:

        messages.error(
            request,
            "Semester Grade Sheet is not available."
        )

        return redirect(
            "student_results"
        )

    return semester_grade_sheet(
        request,
        student.id,
        publication.examination.id,
        semester,
    )

@login_required
def student_combined_grade_sheet(request):

    student = Student.objects.get(
        user=request.user
    )

    latest_publication = (
        ResultPublication.objects.filter(
            examination__combined_results__student=student,
            is_published=True,
        )
        .order_by(
            "-published_on",
            "-id",
        )
        .first()
    )

    if latest_publication is None:

        messages.error(
            request,
            "Combined Grade Sheet is not available."
        )

        return redirect(
            "student_results"
        )

    latest_result = (
        CombinedResult.objects.filter(
            student=student,
            examination=latest_publication.examination,
        )
        .first()
    )

    if latest_result is None:

        messages.error(
            request,
            "Combined Grade Sheet is not available."
        )

        return redirect(
            "student_results"
        )

    return combined_grade_sheet(
        request,
        student.id,
        latest_result.examination.id,
    )

def logout_view(request):
    logout(request)
    return redirect("login")

@login_required
def student_backlog_registration(request):

    examinations = (
        Examination.objects.filter(
            examination_type__code="BLG",
            status="REGISTRATION_OPEN",
            is_active=True,
        ).order_by("-id")
    )

    failed_subjects = []

    selected_examination_id = None

    rejected_registration = None

    if request.method == "POST":

        print(request.POST)

        selected_examination_id = request.POST.get(
            "examination"
        )

        if not selected_examination_id:

            messages.warning(
                request,
                "No backlog examination registration is currently open."
            )

            return render(
                request,
                "accounts/student_backlog_registration.html",
                {
                    "examinations": examinations,
                    "failed_subjects": [],
                    "selected_examination_id": None,
                },
            )

        selected_examination = Examination.objects.get(
            id=selected_examination_id
        )

        selected_examination_id = selected_examination.id

        student = Student.objects.get(
            user=request.user
        )

        rejected_registration = (
            BacklogRegistration.objects.filter(
                student=student,
                examination=selected_examination,
                registration_status="REJECTED",
            )
            .first()
        )

        action = request.POST.get("action")

        if action == "search":

            failed_subjects = get_backlog_subjects(
                student,
                selected_examination,
            )

        elif action == "register":

            selected_subject_ids = request.POST.getlist("subjects")

            selected_subject_ids = request.POST.getlist(
                "subjects"
            )

            if not selected_subject_ids:

                messages.error(
                    request,
                    "Please select at least one subject."
                )

                failed_subjects = get_backlog_subjects(
                    student,
                    selected_examination,
                )

            else:

                existing_registration = (
                    BacklogRegistration.objects.filter(
                        student=student,
                        examination=selected_examination,
                    ).first()
                )

                if existing_registration and (
                    existing_registration.registration_status == "APPROVED"
                ):

                    messages.warning(
                        request,
                        "Your backlog registration has already been approved and is locked."
                    )

                    return redirect(
                        "student_backlog_registration"
                    )

                was_rejected = False

                if existing_registration and (
                    existing_registration.registration_status == "REJECTED"
                ):

                    was_rejected = True

                    existing_registration.registration_status = "SUBMITTED"

                    existing_registration.approved_by = None

                    existing_registration.approved_on = None

                    existing_registration.save(
                        update_fields=[
                            "registration_status",
                            "approved_by",
                            "approved_on",
                        ]
                    )

                    BacklogRegistrationSubject.objects.filter(
                        backlog_registration=existing_registration
                    ).delete()

                backlog_registration, created = (
                    BacklogRegistration.objects.get_or_create(

                        student=student,

                        examination=selected_examination,

                        defaults={
                            "registration_status": "SUBMITTED",
                        },

                    )
                )

                newly_registered = 0

                already_registered = 0

                for subject_id in selected_subject_ids:

                    _, created = (
                        BacklogRegistrationSubject.objects.get_or_create(

                            backlog_registration=backlog_registration,

                            subject_id=subject_id,

                        )
                    )

                    if created:

                        newly_registered += 1

                    else:

                        already_registered += 1

                if newly_registered > 0 and already_registered == 0:

                    messages.success(
                        request,
                        f"{newly_registered} subject(s) registered successfully."
                    )

                elif newly_registered > 0 and already_registered > 0:

                    messages.success(
                        request,
                        f"{newly_registered} subject(s) registered successfully. "
                        f"{already_registered} subject(s) were already registered."
                    )

                else:

                    messages.warning(
                        request,
                        "You have already registered for all the selected subjects."
                    )

                return redirect(
                    "student_backlog_registration"
                )

    return render(
        request,
        "accounts/student_backlog_registration.html",
        {
            "examinations": examinations,
            "failed_subjects": failed_subjects,
            "selected_examination_id": selected_examination_id,
            "rejected_registration": rejected_registration,
        },
    )

@login_required
def coe_backlog_registrations(request):

    if request.method == "POST":

        registration_id = request.POST.get(
            "registration_id"
        )

        action = request.POST.get(
            "action"
        )

        registration = get_object_or_404(
            BacklogRegistration,
            id=registration_id,
            registration_status="SUBMITTED",
        )

        if action == "approve":

            registration.registration_status = "APPROVED"

            registration.approved_by = request.user

            registration.approved_on = timezone.now()

            registration.remarks = ""

            registration.save(
                update_fields=[
                    "registration_status",
                    "approved_by",
                    "approved_on",
                    "remarks",
                ]
            )

            messages.success(
                request,
                "Backlog registration approved successfully."
            )

        elif action == "reject":

            remarks = request.POST.get(
                "remarks",
                ""
            ).strip()

            registration.registration_status = "REJECTED"

            registration.approved_by = request.user

            registration.approved_on = timezone.now()

            registration.remarks = remarks

            registration.save(
                update_fields=[
                    "registration_status",
                    "approved_by",
                    "approved_on",
                    "remarks",
                ]
            )

            messages.warning(
                request,
                "Backlog registration rejected."
            )

        return redirect(
            "coe_backlog_registrations"
        )

    registrations = (
        BacklogRegistration.objects.filter(
            registration_status="SUBMITTED"
        )
        .select_related(
            "student",
            "examination",
        )
        .prefetch_related(
            "subjects__subject",
        )
        .order_by(
            "student__roll_number",
        )
    )

    return render(
        request,
        "accounts/coe_backlog_registrations.html",
        {
            "registrations": registrations,
        },
    )
